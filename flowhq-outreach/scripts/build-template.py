#!/usr/bin/env python3
"""Generates templates/FlowHQ-Outreach-Template.xlsx from lib/schema.js.
Upload it to Google Drive and open it as a Google Sheet (File > Save as Google Sheets)."""
import json, subprocess, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
schema = json.loads(subprocess.check_output(['node', '-e', "const s=require('./lib/schema');console.log(JSON.stringify(s))"], cwd=ROOT))
wb = Workbook()
readme = wb.active
readme.title = 'README'
lines = [
    'FlowHQ Outreach workspace',
    '',
    'Leads_Source: paste or link your raw leads here (any column names; mapping is configurable). The workflows only READ this tab.',
    'Leads: working table written by the workflows. Review drafts here and set approval_status to APPROVED or REJECTED.',
    'Settings: runtime switches. KILL_SWITCH=TRUE stops all sending. SENDING_ENABLED must be TRUE here AND in the workflow config.',
    'Suppression: emails or domains that must never be contacted (value + type email|domain).',
    'Research_Cache, Send_Log, Inbox_Log, Error_Log: written by the workflows; do not edit headers.',
    '',
    'Do not rename tabs or header cells. Column lead_id must stay in column A of Leads.',
]
for i, l in enumerate(lines, 1):
    readme.cell(row=i, column=1, value=l)
readme['A1'].font = Font(bold=True, size=14)
readme.column_dimensions['A'].width = 140

head_font = Font(bold=True, color='FFFFFF')
head_fill = PatternFill('solid', fgColor='1F3A5F')
tabs = {'Leads_Source': ['lead_id', 'business_name', 'contact_name', 'email', 'website', 'country', 'industry', 'source', 'notes', 'outreach_status']}
tabs.update(schema['TABS'])
for name, cols in tabs.items():
    ws = wb.create_sheet(name)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(vertical='center')
        ws.column_dimensions[cell.column_letter].width = max(12, min(40, len(c) + 4))
    ws.freeze_panes = 'B2' if name == 'Leads' else 'A2'
    if name == 'Settings':
        for i, r in enumerate(schema['DEFAULT_SETTINGS_ROWS'], 2):
            ws.cell(row=i, column=1, value=r['key']); ws.cell(row=i, column=2, value=r['value']); ws.cell(row=i, column=3, value=r['notes'])
        ws.column_dimensions['C'].width = 110
    if name == 'Leads':
        col = cols.index('approval_status') + 1
        letter = ws.cell(row=1, column=col).column_letter
        dv = DataValidation(type='list', formula1='"PENDING,APPROVED,REJECTED"', allow_blank=True)
        ws.add_data_validation(dv); dv.add(f'{letter}2:{letter}5000')
        for k in ('email_body', 'original_row_json', 'signals_json', 'validation_issues', 'pages_reviewed'):
            ws.column_dimensions[ws.cell(row=1, column=cols.index(k) + 1).column_letter].width = 60
os.makedirs(os.path.join(ROOT, 'templates'), exist_ok=True)
out = os.path.join(ROOT, 'templates', 'FlowHQ-Outreach-Template.xlsx')
wb.save(out)
print('wrote', out, [ws.title for ws in wb.worksheets])
